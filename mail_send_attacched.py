import base64
import mimetypes
import os
from email.message import EmailMessage
from google.auth.transport.requests import Request
from google.oauth2.credentials import Credentials
from google_auth_oauthlib.flow import InstalledAppFlow
from googleapiclient.discovery import build
from googleapiclient.errors import HttpError
from openpyxl import load_workbook
from docx import Document

# Gmail API scope
SCOPES = ['https://www.googleapis.com/auth/gmail.send']

def authentication():
    credentials = None

    if os.path.exists('token.json'):
        credentials = Credentials.from_authorized_user_file('token.json', SCOPES)

    if not credentials or not credentials.valid:
        if credentials and credentials.expired and credentials.refresh_token:
            credentials.refresh(Request())
        else:
            flow = InstalledAppFlow.from_client_secrets_file('E:/PRACTICE/NumPy/client_secret_apps.json', SCOPES)
            credentials = flow.run_local_server(port=0)

        with open('token.json', 'w') as token:
            token.write(credentials.to_json())

    return credentials

def read_email_addresses_from_excel(excel_file, sheet_name, email_column):
    workbook = load_workbook(excel_file)
    sheet = workbook[sheet_name]
    
    email_addresses = [cell.value for cell in sheet[email_column] if cell.value is not None]
    
    return email_addresses

def read_docx_content(file_path):
    try:
        doc = Document(file_path)
        paragraphs = [paragraph.text for paragraph in doc.paragraphs]
        return '\n'.join(paragraphs)
    except Exception as e:
        print(f"Error reading Word document: {e}")
        return None

def create_message(sender, to, subject, body):
    mime_message = EmailMessage()
    mime_message['From'] = sender
    mime_message['To'] = to
    mime_message['Subject'] = subject
    mime_message.set_content(body)
    
    return {'raw': base64.urlsafe_b64encode(mime_message.as_bytes()).decode()}

def send_message(service, user_id, message):
    try:
        message = service.users().messages().send(userId=user_id, body=message).execute()
        print(f"Message Id: {message['id']}")
        return message
    except HttpError as error:
        print(f"An error occurred: {error}")

if __name__ == "__main__":
    excel_file = 'E:/PRACTICE/NumPy/data.xlsx'
    sheet_name = 'Gmail'
    email_column = 'A'
    wordpad_file = 'E:/PRACTICE/NumPy/template.docx'

    recipients = read_email_addresses_from_excel(excel_file, sheet_name, email_column)
    wordpad_content = read_docx_content(wordpad_file)

    credentials = authentication()
    service = build(serviceName='gmail', version='v1', credentials=credentials)

    sender_email = 'ben.smith41349@gmail.com'  # replace with your Gmail address
    subject = 'Re-Design Your Web'

    for recipient_email in recipients:
        message = create_message(sender_email, recipient_email, subject, wordpad_content)
        send_message(service, 'me', message)
